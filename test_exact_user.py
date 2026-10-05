import os
import sys
import pptx
import fitz

# Create presentation with exact user text
prs = pptx.Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "تصميم هندسي لمنصة بذور المعرفة: جهاز لوحي فائق الاستقلالية وشبكة تعليمية لامركزية خالية من الوايفاي"
slide.placeholders[1].text = "السياق الاستراتيجي وأهداف التحدي التكنولوجي (Technology E3)\nيستشعر هذا التقرير البدء في العمل صمماً هندسياً وتقنياً متكاملاً لتطوير استجابة للتحدي الثاني لشركة تكنولوجيا E3 ضمن هاكاثون 2026."

temp_pptx = os.path.abspath("test_exact_user.pptx")
prs.save(temp_pptx)

from app.services.pptx_converter import convert_pptx_bytes_to_pdf_bytes

with open(temp_pptx, "rb") as f:
    data = f.read()

pdf_bytes = convert_pptx_bytes_to_pdf_bytes(data)

doc = fitz.open(stream=pdf_bytes, filetype="pdf")
print("Generated PDF page count:", len(doc))
for i, page in enumerate(doc):
    print(f"--- Page {i+1} Text ---")
    print(page.get_text()[:400])

pix = doc[0].get_pixmap(dpi=150)
pix.save("test_converted_exact.png")
print("Saved test_converted_exact.png successfully!")

if os.path.exists(temp_pptx):
    os.remove(temp_pptx)
